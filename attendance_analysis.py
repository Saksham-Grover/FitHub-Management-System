"""
attendance_analysis.py - Attendance statistics report and chart (Member 3).

export_attendance_report()   -> Excel file with attendance by class type and by trainer
show_attendance_chart(parent) -> "Attendance per Week" chart embedded in a Tkinter window

Only bookings with BookingStatus = 'Booked' are counted (cancelled bookings are
excluded). The attendance rate is Attended / (Attended + No-show), so bookings
still marked 'Pending' do not lower the rate.
"""
import tkinter as tk
from datetime import datetime
from tkinter import filedialog, messagebox

import mysql.connector
from matplotlib.backends.backend_tkagg import FigureCanvasTkAgg
from matplotlib.figure import Figure
from matplotlib.ticker import MaxNLocator
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from db_config import get_connection

STATS_SQL = """
    SELECT {label},
           COUNT(*),
           SUM(b.AttendanceStatus = 'Attended'),
           SUM(b.AttendanceStatus = 'No-show'),
           SUM(b.AttendanceStatus = 'Pending'),
           ROUND(100 * SUM(b.AttendanceStatus = 'Attended')
                 / NULLIF(SUM(b.AttendanceStatus IN ('Attended', 'No-show')), 0), 1)
    FROM Booking b
    JOIN Class c ON b.ClassID = c.ClassID
    {joins}
    WHERE b.BookingStatus = 'Booked'
    GROUP BY {label}
    ORDER BY {label}
"""

WEEKLY_SQL = """
    SELECT DATE_SUB(b.ClassDate, INTERVAL WEEKDAY(b.ClassDate) DAY) AS week_start,
           SUM(b.AttendanceStatus = 'Attended'),
           SUM(b.AttendanceStatus = 'No-show')
    FROM Booking b
    WHERE b.BookingStatus = 'Booked'
      AND b.AttendanceStatus IN ('Attended', 'No-show')
    GROUP BY week_start
    ORDER BY week_start
"""

HEADERS = ["Total Bookings", "Attended", "No-shows", "Pending", "Attendance Rate (%)"]


def fetch_rows(sql):
    """Run a read-only query and return all rows."""
    connection = None
    cursor = None
    try:
        connection = get_connection()
        cursor = connection.cursor()
        cursor.execute(sql)
        return cursor.fetchall()
    finally:
        if cursor:
            cursor.close()
        if connection:
            connection.close()


def clean_row(row):
    """Turn MySQL Decimal values into plain ints/floats so Excel stores numbers."""
    label, total, attended, no_shows, pending, rate = row
    return [label, int(total), int(attended or 0), int(no_shows or 0),
            int(pending or 0), float(rate) if rate is not None else None]


def add_sheet(workbook, title, first_header, rows):
    """Write one formatted worksheet."""
    sheet = workbook.create_sheet(title)
    sheet.append([first_header] + HEADERS)
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="2C3E50")
        cell.alignment = Alignment(horizontal="center")
    for row in rows:
        sheet.append(clean_row(row))
    for index in range(1, len(HEADERS) + 2):
        sheet.column_dimensions[get_column_letter(index)].width = 22


def export_attendance_report():
    """Ask where to save, then write the attendance statistics to an Excel file."""
    try:
        by_type = fetch_rows(STATS_SQL.format(label="c.ClassType", joins=""))
        by_trainer = fetch_rows(STATS_SQL.format(
            label="t.TrainerName", joins="JOIN Trainer t ON c.TrainerID = t.TrainerID"))
    except mysql.connector.Error as error:
        messagebox.showerror("Database Error", str(error))
        return

    if not by_type:
        messagebox.showinfo("No Data", "There are no bookings to report on yet.")
        return

    path = filedialog.asksaveasfilename(
        title="Save attendance report",
        defaultextension=".xlsx",
        initialfile=f"attendance_statistics_{datetime.now():%Y%m%d_%H%M}.xlsx",
        filetypes=[("Excel workbook", "*.xlsx")])
    if not path:
        return  # the user cancelled

    workbook = Workbook()
    workbook.remove(workbook.active)
    add_sheet(workbook, "By Class Type", "Class Type", by_type)
    add_sheet(workbook, "By Trainer", "Trainer", by_trainer)

    try:
        workbook.save(path)
    except PermissionError:
        messagebox.showerror("Cannot Save", "The file is open in Excel. Close it and try again.")
        return
    messagebox.showinfo("Report Saved", f"Attendance report saved to:\n{path}")


def show_attendance_chart(parent):
    """Open a window with the Attendance per Week bar chart."""
    try:
        rows = fetch_rows(WEEKLY_SQL)
    except mysql.connector.Error as error:
        messagebox.showerror("Database Error", str(error))
        return
    if not rows:
        messagebox.showinfo("No Data", "There is no attendance data to chart yet.")
        return

    labels = [row[0].strftime("%d %b") for row in rows]
    attended = [int(row[1] or 0) for row in rows]
    no_shows = [int(row[2] or 0) for row in rows]
    positions = list(range(len(rows)))

    window = tk.Toplevel(parent)
    window.title("FitHub - Attendance per Week")
    window.geometry("900x550")

    figure = Figure(figsize=(9, 5), dpi=100)
    axes = figure.add_subplot(111)
    axes.bar([p - 0.2 for p in positions], attended, width=0.4, label="Attended", color="#27ae60")
    axes.bar([p + 0.2 for p in positions], no_shows, width=0.4, label="No-show", color="#c0392b")
    axes.set_xticks(positions)
    axes.set_xticklabels(labels)
    axes.set_title("Attendance per Week")
    axes.set_xlabel("Week starting (Monday)")
    axes.set_ylabel("Number of bookings")
    axes.yaxis.set_major_locator(MaxNLocator(integer=True))
    axes.legend()
    figure.tight_layout()

    canvas = FigureCanvasTkAgg(figure, master=window)
    canvas.draw()
    canvas.get_tk_widget().pack(fill="both", expand=True)
